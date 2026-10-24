# Tests del sidecar de herramientas: documentos, PII y NO-BLOQUEO del lector.
# Ejecutar:  python-runtime/python.exe -m pytest python/tools_sidecar/tests -q
import json
import os
import subprocess
import sys
import threading
import time

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
SIDECAR_DIR = os.path.join(ROOT, "python", "tools_sidecar")
PYTHON = sys.executable

sys.path.insert(0, SIDECAR_DIR)

import documents  # noqa: E402
import main  # noqa: E402
import privacy  # noqa: E402


def test_redacta_etiqueta_de_hoja_sin_devolver_correo_en_metadata():
    documento = {"markdown": "Hoja juan@empresa.mx", "tables": [
        {"sheet": "juan@empresa.mx", "rows": [["juan@empresa.mx"]]}
    ], "metadata": {"sheets": 1}}
    resultado = main._redact_document(documento, privacy)
    assert resultado["tables"][0]["sheet"] == "[EMAIL]"
    assert "juan@empresa.mx" not in json.dumps(resultado)


# ── Documentos ───────────────────────────────────────────────────────────────

def _build_xlsx(path: str) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ventas"
    sheet.append(["Producto", "Unidades", "Total"])
    sheet.append(["Teclado", 3, 1500])
    sheet.append(["Monitor", 1, 4200])
    workbook.save(path)


def test_xlsx_se_convierte_a_markdown_con_tabla(tmp_path):
    archivo = str(tmp_path / "ventas.xlsx")
    _build_xlsx(archivo)

    resultado = documents.parse_document(archivo)

    assert "## Hoja: Ventas" in resultado["markdown"]
    assert "| Producto | Unidades | Total |" in resultado["markdown"]
    assert "Monitor" in resultado["markdown"]
    # La tabla tambien viaja estructurada (no solo como texto).
    assert resultado["tables"][0]["rows"][0] == ["Producto", "Unidades", "Total"]
    assert resultado["metadata"]["format"] == "xlsx"


def test_docx_conserva_encabezados_y_tablas(tmp_path):
    from docx import Document

    archivo = str(tmp_path / "informe.docx")
    documento = Document()
    documento.add_heading("Resumen", level=1)
    documento.add_paragraph("Contenido del informe.")
    tabla = documento.add_table(rows=2, cols=2)
    tabla.cell(0, 0).text = "Concepto"
    tabla.cell(0, 1).text = "Monto"
    tabla.cell(1, 0).text = "Licencias"
    tabla.cell(1, 1).text = "12000"
    documento.save(archivo)

    resultado = documents.parse_document(archivo)

    assert "## Resumen" in resultado["markdown"]
    # mammoth (TS) aplanaba las tablas; aqui se conservan.
    assert "| Concepto | Monto |" in resultado["markdown"]
    assert resultado["tables"][0]["rows"][1] == ["Licencias", "12000"]


def test_formato_no_soportado_da_error_con_codigo(tmp_path):
    archivo = tmp_path / "notas.txt"
    archivo.write_text("hola", encoding="utf-8")

    try:
        documents.parse_document(str(archivo))
        assert False, "deberia haber lanzado DocumentError"
    except documents.DocumentError as error:
        assert error.code == "DOC_UNSUPPORTED"


def test_archivo_inexistente_da_error_con_codigo():
    try:
        documents.parse_document("C:/no/existe.pdf")
        assert False, "deberia haber lanzado DocumentError"
    except documents.DocumentError as error:
        assert error.code == "DOC_NOT_FOUND"


def test_pptx_preserva_texto_y_tablas(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Resumen ficticio"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
    table.cell(0, 0).text = "Concepto"
    table.cell(0, 1).text = "Monto"
    table.cell(1, 0).text = "Prueba"
    table.cell(1, 1).text = "120"
    file = tmp_path / "ficticio.pptx"
    presentation.save(str(file))
    result = documents.parse_document(str(file))
    assert "Resumen ficticio" in result["markdown"]
    assert result["tables"][0]["rows"][1] == ["Prueba", "120"]


def _write_pdf(path):
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R 5 0 R] /Count 2 >>"]
    for page, stream in [(3, 4), (5, 6)]:
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {stream} 0 R /Resources << /Font << /F1 7 0 R >> >> >>".encode())
        text = f"BT /F1 12 Tf 72 720 Td (Pagina ficticia {page}) Tj ET".encode()
        objects.append(b"<< /Length " + str(len(text)).encode() + b" >>\nstream\n" + text + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    content = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, item in enumerate(objects, 1):
        offsets.append(len(content))
        content.extend(f"{index} 0 obj\n".encode() + item + b"\nendobj\n")
    start = len(content)
    content.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        content.extend(f"{offset:010} 00000 n \n".encode())
    content.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n".encode())
    path.write_bytes(content)


def test_pdf_declara_parcialidad_sin_inventar_total(tmp_path):
    file = tmp_path / "ficticio.pdf"
    _write_pdf(file)
    partial = documents.parse_document(str(file), 1)
    assert partial["metadata"]["pagesParsed"] == 1
    assert partial["metadata"]["truncated"] is True
    assert "pages" not in partial["metadata"]
    complete = documents.parse_document(str(file), 2)
    assert complete["metadata"]["pages"] == 2
    assert "Pagina ficticia 5" in complete["markdown"]


def test_xlsm_conserva_lectura_sin_ejecutar_macros(tmp_path):
    file = tmp_path / "ficticio.xlsm"
    _build_xlsx(str(file))
    result = documents.parse_document(str(file))
    assert result["metadata"]["format"] == "xlsm"
    assert "Monitor" in result["markdown"]


def test_excel_dimension_falsa_no_oculta_celdas_reales(tmp_path):
    import re
    import zipfile
    file = tmp_path / "original.xlsx"
    changed = tmp_path / "dimension-ficticia.xlsx"
    _build_xlsx(str(file))
    with zipfile.ZipFile(file) as source, zipfile.ZipFile(changed, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', content)
            target.writestr(item.filename, content)
    result = documents.parse_document(str(changed))
    assert "Monitor" in result["markdown"]
    assert result["tables"][0]["rows"][2] == ["Monitor", "1", "4200"]


def test_excel_conserva_espacios_interiores_y_saltos_en_celdas(tmp_path):
    from openpyxl import Workbook
    file = tmp_path / "espacios-ficticios.xlsx"
    book = Workbook()
    book.active.append(["  uno  dos\ntres  "])
    book.save(file)
    book.close()
    result = documents.parse_document(str(file))
    assert result["tables"][0]["rows"][0][0] == "uno  dos\ntres"
    assert "| uno  dos tres |" in result["markdown"]


# ── PII ──────────────────────────────────────────────────────────────────────

def test_redacta_pii_mexicana_antes_de_la_nube():
    texto = (
        "Cliente Juan con RFC GODE561231GR8, CURP GODE561231HDFSRL05, "
        "tarjeta 4111 1111 1111 1111, tel 55 1234 5678 y correo juan@empresa.mx."
    )

    resultado = privacy.redact(texto)

    assert "[RFC]" in resultado["redacted"]
    assert "[CURP]" in resultado["redacted"]
    assert "[TARJETA]" in resultado["redacted"]
    assert "[EMAIL]" in resultado["redacted"]
    assert "[TELEFONO]" in resultado["redacted"]
    # El dato sensible ya no aparece en el texto que se enviaria al modelo.
    assert "GODE561231GR8" not in resultado["redacted"]
    assert "4111" not in resultado["redacted"]
    assert "juan@empresa.mx" not in resultado["redacted"]
    # Y tampoco se devuelve en el detalle de entidades.
    assert all("text" not in entidad for entidad in resultado["entities"])


def test_no_marca_como_tarjeta_un_numero_que_no_pasa_luhn():
    # Un folio de 16 digitos no valido por Luhn NO debe redactarse como tarjeta.
    resultado = privacy.redact("Folio 1234567812345678 del pedido.")
    assert "[TARJETA]" not in resultado["redacted"]
    assert "1234567812345678" in resultado["redacted"]


def test_texto_sin_pii_queda_intacto():
    texto = "La reunion es el martes a las diez en la sala grande."
    resultado = privacy.redact(texto)
    assert resultado["redacted"] == texto
    assert resultado["redactedCount"] == 0


def test_parse_document_con_redact_no_devuelve_pii_ni_en_las_tablas(tmp_path):
    """Con redact=true la PII se elimina DENTRO del sidecar: nunca llega a Electron."""
    from openpyxl import Workbook

    archivo = str(tmp_path / "clientes.xlsx")
    workbook = Workbook()
    hoja = workbook.active
    hoja.title = "Clientes"
    hoja.append(["Nombre", "RFC", "Correo"])
    hoja.append(["Juan", "GODE561231GR8", "juan@empresa.mx"])
    workbook.save(archivo)

    resultado = main._redact_document(documents.parse_document(archivo), privacy)

    assert "GODE561231GR8" not in resultado["markdown"]
    assert "juan@empresa.mx" not in resultado["markdown"]
    assert "[RFC]" in resultado["markdown"]
    # Las celdas de las tablas estructuradas tambien van redactadas.
    celdas = [celda for fila in resultado["tables"][0]["rows"] for celda in fila]
    assert "GODE561231GR8" not in celdas
    assert "[RFC]" in celdas
    assert resultado["metadata"]["piiRedacted"] >= 2
    assert resultado["metadata"]["piiEngine"] == privacy.ENGINE


# ── El lector NUNCA se bloquea (el defecto que tiene el sidecar de voz) ──────

def test_una_tarea_larga_no_bloquea_los_demas_comandos(tmp_path):
    """Con un XLSX grande en curso, `ping` debe responder igualmente rapido."""
    grande = str(tmp_path / "grande.xlsx")
    from openpyxl import Workbook

    workbook = Workbook()
    hoja = workbook.active
    for fila in range(1500):
        hoja.append([f"dato-{fila}-{columna}" for columna in range(12)])
    workbook.save(grande)

    proceso = subprocess.Popen(
        [PYTHON, "-u", os.path.join(SIDECAR_DIR, "main.py")],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
        text=True, encoding="utf-8", cwd=SIDECAR_DIR,
    )
    try:
        respuestas: dict = {}
        errores = []
        listo = threading.Event()

        def leer():
            for linea in proceso.stdout:
                try:
                    mensaje = json.loads(linea)
                except json.JSONDecodeError:
                    continue
                if mensaje.get("event") == "ready":
                    listo.set()
                elif mensaje.get("event") == "error":
                    errores.append(mensaje.get("message"))
                elif mensaje.get("id") is not None:
                    respuestas[mensaje["id"]] = (mensaje, time.monotonic())

        threading.Thread(target=leer, daemon=True).start()
        assert listo.wait(timeout=15), "el sidecar no arranco"

        # 1) Se lanza el trabajo pesado.  2) Inmediatamente despues, un ping.
        inicio = time.monotonic()
        proceso.stdin.write(json.dumps({"id": 1, "cmd": "parse_document", "params": {"file_path": grande}}) + "\n")
        proceso.stdin.write(json.dumps({"id": 2, "cmd": "ping"}) + "\n")
        proceso.stdin.flush()

        limite = time.monotonic() + 30
        while 2 not in respuestas and time.monotonic() < limite:
            time.sleep(0.02)

        assert 2 in respuestas, "el ping nunca respondio (el lector estaba bloqueado)"
        ping_ms = (respuestas[2][1] - inicio) * 1000
        assert respuestas[2][0]["ok"] is True
        # Si el comando pesado corriera en el hilo lector, el ping tardaria segundos.
        assert ping_ms < 1000, f"el ping tardo {ping_ms:.0f} ms: el lector se bloqueo"

        while 1 not in respuestas and time.monotonic() < limite:
            time.sleep(0.05)
        if 1 not in respuestas:
            proceso.kill()
            proceso.wait(timeout=5)
        assert 1 in respuestas, f"la lectura del documento no respondió: {errores}"
        assert respuestas[1][0]["ok"] is True
        assert respuestas[1][0]["data"]["metadata"]["format"] == "xlsx"

        # Un rechazo presupuestado no impide atender el siguiente comando.
        import zipfile
        limitado = str(tmp_path / "limite.docx")
        with zipfile.ZipFile(limitado, "w") as archive:
            archive.writestr("word/document.xml", '<document><gridSpan val="2147483647"/></document>')
        proceso.stdin.write(json.dumps({"id": 3, "cmd": "parse_document", "params": {"file_path": limitado, "redact": True}}) + "\n")
        proceso.stdin.write(json.dumps({"id": 4, "cmd": "ping"}) + "\n")
        proceso.stdin.flush()
        limite = time.monotonic() + 10
        while (3 not in respuestas or 4 not in respuestas) and time.monotonic() < limite:
            time.sleep(0.02)
        assert respuestas[3][0]["error"]["code"] == "DOC_RESOURCE_LIMIT"
        assert limitado not in json.dumps(respuestas[3][0])
        assert respuestas[4][0]["ok"] is True

        proceso.stdin.write(json.dumps({"id": 9, "cmd": "shutdown"}) + "\n")
        proceso.stdin.flush()
        proceso.wait(timeout=5)

    finally:
        if proceso.poll() is None:
            proceso.kill()
            proceso.wait(timeout=5)
        proceso.stdin.close()
        proceso.stdout.close()
