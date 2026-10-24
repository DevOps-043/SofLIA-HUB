"""Lectura local con admisión y construcción presupuestadas; sin extraer archivos."""
import os
import stat
from document_errors import DocumentError, resource_limit
from document_budget import (
    DocumentBuilder, MAX_FILE_BYTES, MAX_PDF_PAGES, MAX_SHEET_ROWS,
    MAX_MARKDOWN_CHARS, MAX_SECTIONS, validate_result,
)
from office_preflight import validate_office

SUPPORTED_EXTENSIONS = {".pdf", ".xlsx", ".xlsm", ".pptx", ".docx"}


def parse_document(file_path: str, max_pages: int = MAX_PDF_PAGES) -> dict:
    """Conserva la API por ruta; bibliotecas y preflight usan el mismo descriptor."""
    if not isinstance(file_path, str) or not file_path or type(max_pages) is not int or not 1 <= max_pages <= MAX_PDF_PAGES:
        raise DocumentError("TOOLS_INVALID_INPUT", "La solicitud de lectura no es válida.")
    extension = os.path.splitext(file_path)[1].lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentError("DOC_UNSUPPORTED", "El formato del documento no está admitido.")
    try:
        before = os.stat(file_path, follow_symlinks=False)
        if not stat.S_ISREG(before.st_mode):
            raise DocumentError("DOC_PARSE_FAILED", "El documento no es un archivo regular local.")
        if before.st_size > MAX_FILE_BYTES:
            raise DocumentError("DOC_TOO_LARGE", "El documento supera 50 MiB.")
        flags = os.O_RDONLY | getattr(os, "O_BINARY", 0) | getattr(os, "O_NOFOLLOW", 0) | getattr(os, "O_NONBLOCK", 0)
        with os.fdopen(os.open(file_path, flags), "rb") as source:
            opened = os.fstat(source.fileno())
            if not stat.S_ISREG(opened.st_mode) or opened.st_size > MAX_FILE_BYTES:
                raise DocumentError("DOC_TOO_LARGE", "La fuente abierta no es un documento regular dentro del límite.")
            prior = lambda value: (value.st_dev, value.st_ino, value.st_size, value.st_mtime_ns)
            if prior(opened) != prior(before):
                raise DocumentError("DOCUMENT_CHANGED", "El documento cambió durante la lectura.")
            if extension != ".pdf":
                validate_office(source, opened.st_size)
            if extension == ".pdf":
                result = _parse_pdf(source, max_pages)
            elif extension in (".xlsx", ".xlsm"):
                result = _parse_xlsx(source)
            elif extension == ".pptx":
                result = _parse_pptx(source)
            else:
                result = _parse_docx(source)
            after = os.fstat(source.fileno())
            identity = lambda value: (value.st_dev, value.st_ino, value.st_size, value.st_mtime_ns, value.st_ctime_ns)
            # Windows puede diferir ctime entre stat de ruta y primer fstat.
            if prior(before) != prior(opened) or identity(opened) != identity(after):
                raise DocumentError("DOCUMENT_CHANGED", "El documento cambió durante la lectura.")
    except FileNotFoundError as error:
        raise DocumentError("DOC_NOT_FOUND", "No se encontró el documento seleccionado.") from error
    result["metadata"]["format"] = extension.lstrip(".")
    result["metadata"]["sizeBytes"] = opened.st_size
    return validate_result(result)


def _parse_pdf(source, max_pages):
    from reader_dependencies import assert_pdf_dependency
    assert_pdf_dependency()
    import pdfplumber
    builder = DocumentBuilder()
    pages_with_text = 0
    try:
        # La página adicional permite declarar parcialidad sin inventar el total.
        with pdfplumber.open(source, pages=list(range(1, max_pages + 2))) as pdf:
            pages = pdf.pages
            parsed = min(len(pages), max_pages)
            for index, page in enumerate(pages[:max_pages], start=1):
                builder.text(f"## Pagina {index}")
                text = (page.extract_text() or "").strip()
                if text:
                    pages_with_text += 1
                    builder.text(text)
                for table_index, rows in enumerate(page.extract_tables() or [], start=1):
                    builder.table(rows, page=index, index=table_index)
                builder.text("")
            metadata = {"pagesParsed": parsed}
            if len(pages) > max_pages:
                metadata["truncated"] = True
            else:
                metadata["pages"] = len(pages)
    except DocumentError:
        raise
    except Exception as error:
        code = "DOC_ENCRYPTED" if any(word in str(error).lower() for word in ("password", "encrypt")) else "DOC_PARSE_FAILED"
        raise DocumentError(code, "No se pudo interpretar el PDF.") from error
    if not pages_with_text and not builder.tables:
        raise DocumentError("DOC_SCANNED", "El PDF no tiene texto seleccionable. Requiere OCR.", recoverable=True)
    return builder.finish(metadata)


def _parse_xlsx(source):
    from openpyxl import load_workbook
    workbook = None
    builder = DocumentBuilder()
    try:
        workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
        if len(workbook.worksheets) > MAX_SECTIONS:
            raise resource_limit()
        for sheet in workbook.worksheets:
            builder.text(f"## Hoja: {sheet.title}")
            # La geometría real ya pasó preflight; no confiar en dimension XML.
            sheet.reset_dimensions()
            tables_before = len(builder.tables)
            builder.table(sheet.iter_rows(values_only=True), normalize_whitespace=False, sheet=sheet.title)
            if len(builder.tables) == tables_before:
                builder.text("_(hoja vacia)_")
            builder.text("")
        return builder.finish({"sheets": len(workbook.worksheets)})
    except DocumentError:
        raise
    except Exception as error:
        raise DocumentError("DOC_PARSE_FAILED", "No se pudo interpretar el Excel.") from error
    finally:
        if workbook is not None:
            workbook.close()


def _parse_pptx(source):
    from pptx import Presentation
    builder = DocumentBuilder()
    try:
        presentation = Presentation(source)
        if len(presentation.slides) > MAX_SECTIONS:
            raise resource_limit()
        for index, slide in enumerate(presentation.slides, start=1):
            builder.text(f"## Diapositiva {index}")
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    builder.text(shape.text_frame.text.strip())
                if getattr(shape, "has_table", False):
                    builder.table(([cell.text for cell in row.cells] for row in shape.table.rows), slide=index)
            builder.text("")
        return builder.finish({"slides": len(presentation.slides)})
    except DocumentError:
        raise
    except Exception as error:
        raise DocumentError("DOC_PARSE_FAILED", "No se pudo interpretar la presentación.") from error


def _parse_docx(source):
    from docx import Document
    builder = DocumentBuilder()
    try:
        document = Document(source)
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").lower()
            if style.startswith("heading"):
                level = "".join(char for char in style if char.isdigit()) or "1"
                text = f"{'#' * min(int(level) + 1, 6)} {text}"
            builder.text(text)
        for index, table in enumerate(document.tables, start=1):
            builder.table(([cell.text for cell in row.cells] for row in table.rows), index=index)
        return builder.finish({"paragraphs": len(document.paragraphs), "tables": len(builder.tables)})
    except DocumentError:
        raise
    except Exception as error:
        raise DocumentError("DOC_PARSE_FAILED", "No se pudo interpretar el documento Word.") from error
